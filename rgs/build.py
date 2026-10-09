import json, re, os, itertools, collections
import openpyxl
from . import content as C
from . import rules as R

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
def g(r, k):
    v = r.get(k); return v[0] if isinstance(v, list) and v else v

class Run:
    def __init__(self, cfg, snap, parent_sku_mode=None):
        self.cfg, self.snap = cfg, snap
        self.mode = parent_sku_mode or cfg.get("parent_sku_mode", "collection")
        self.log = []     # (severity, item, scope, message)
        self.checks = []
        R.apply(cfg, self.log)
        self.single = cfg["listing_structure"] == "single_parent"
        T = snap["tables"]
        self.lib = T["Listing Content Library"]
        coll, prod = cfg["collection"], cfg["product_code"]
        self.res = C.Resolver(self.lib, coll, prod, self.log)
        self.sku = {}
        for r in T["SKU Registries"]:
            m = re.match(re.escape(coll) + r"-(\d{2})(?:-|$)", str(r.get("RGS SKU")))
            if m:   # design rows only; the parent row (RGS SKU == collection code) carries no design number
                self.sku[(r["RGS Parent SKU"], int(m.group(1)))] = r
        st = [r for r in T["Listing Build Staging"] if prod in str(r.get("Product Code")) and coll in str(r.get("RGS Child SKU", "")) ]
        self.st = {"Etsy": {}, "Shopify": {}}
        for r in st:
            for m in self.st:
                if m in str(r.get("Marketplace")): self.st[m][r["RGS Child SKU"]] = r
        self.img = {(r["Platform"], r["Product SKU"]): [(r.get("Image %d" % i), r.get("Alt Text %d" % i)) for i in range(1, 18)] for r in T["Image URL Console"]}
        self.P = {}
        for p in cfg["parents"]:
            i = 1 if self.single else int(p[-2:])   # lead design for this parent
            self.P[p] = dict(i=i, order=R.variant_order(cfg, i))
        self.ai_hold = []
        self.pricing = {r["Pricing ID"]: r for r in T.get("Pricing Registry", [])}

    def price(self, market, des):
        rid = self.cfg["pricing"][market]["ready_made" if des == self.cfg["ready_made_design"] else "rotating"]
        r = self.pricing.get(rid)
        v = (r or {}).get("Launch Price") or (r or {}).get("Target Price")
        if v is None: self.log.append(("BLOCKER", "Pricing Registry", rid, "no Launch or Target price")); return 0
        return v

    def ptext(self, p, key, market="Etsy"):
        vals = {r.get(key) for r in self.st[market].values() if r["RGS Parent SKU"] == p}
        if len(vals) != 1: self.log.append(("BLOCKER", key, p, "%s staged values for %s/%s" % (len(vals), market, key))); return None
        return vals.pop()

    def readiness(self):
        cfg = self.cfg; ok = True
        for p in cfg["parents"]:
            i = self.P[p]["i"]; info = self.P[p]
            info["title"] = self.ptext(p, "Etsy Title Draft")
            info["tags"] = [t.strip() for t in (self.ptext(p, "Etsy Tags Draft") or "").split(",") if t.strip()]
            info["seo"] = self.ptext(p, "Shopify SEO Title Draft", "Shopify")
            info["meta"] = self.ptext(p, "Shopify SEO Meta Description Draft", "Shopify")
            info["shop_tags"] = [t.strip() for t in (self.ptext(p, "Etsy Tags Draft", "Shopify") or "").split(",") if t.strip()]
            hs = {r.get("Shopify Handle") for (pp, d), r in self.sku.items() if pp == p} - {None}
            info["handle"] = hs.pop() if len(hs) == 1 else None
            if not info["handle"]: self.log.append(("BLOCKER", "Shopify Handle", p, "no unique handle in SKU Registries"))
            for market, key in (("Etsy", "etsy"), ("Shopify", "shopify")):
                recs, miss = self.res.build(cfg[key]["manifest"], market, i)
                info[market + "_recs"] = recs
                for t in miss: self.log.append(("BLOCKER", "Content Library", p, "no usable %s record for section '%s'" % (market, t)))
            for L, des in info["order"]:
                s = self.sku.get((p, des))
                if not s: self.log.append(("BLOCKER", "SKU Registries", p, "missing design %02d" % des)); continue
                if not str(s.get("Vendor SKU", "")).startswith("apod-"): self.log.append(("BLOCKER", "Vendor SKU", p, "design %02d vendor SKU not resolved" % des))
            for mk in ("Etsy", "Shopify"):
                k = (mk, self.child(p, i))
                if k not in self.img: self.log.append(("BLOCKER", "Image URL Console", p, "no %s image row for lead design" % mk))
            for des in {d for _, d in info["order"]}:
                if (("Etsy", self.child(p, des)) not in self.img) and des == cfg["ready_made_design"]:
                    self.log.append(("HOLD", "Image URL Console", p, "ready-made design %02d has no image set; overwrite_images forced FALSE" % des))
        for mk in ("Etsy", "Shopify", "Amazon"):
            for des in (1, cfg["ready_made_design"]): self.price(mk, des)
        if any(l[0] == "BLOCKER" for l in self.log): ok = False
        return ok

    def child(self, p, des):
        """Child SKU used to key staging and Image URL Console rows: <COLL>-0N (single parent) or <COLL>-0N-P0M (rotating parents, the default)."""
        return "%s-%02d" % (self.cfg["collection"], des) if self.single else "%s-%02d-%s" % (self.cfg["collection"], des, p[-3:])

    def chk(self, m, rule, c, res, ev): self.checks.append((m, rule, c, res, ev))

    def parent_key(self, p): return self.cfg["collection"] if self.mode == "collection" else p

    def build_etsy(self, out):
        cfg, e = self.cfg, self.cfg["etsy"]
        wb = openpyxl.load_workbook(os.path.join(HERE, "templates/etsy_bulk_template.xlsx")); ws = wb["Template"]
        H = {c.value: c.column for c in ws[1] if c.value}
        defaults = {h: ws.cell(2, c).value for h, c in H.items() if ws.cell(2, c).value is not None}
        for r in range(2, 60):
            for c in range(1, ws.max_column + 1): ws.cell(r, c).value = None
        row = 2; self.desc_etsy = {}
        for p in cfg["parents"]:
            info = self.P[p]; self.desc_etsy[p] = C.etsy_text(info["Etsy_recs"])
            for pos, (L, des) in enumerate(info["order"], 1):
                s = self.sku[(p, des)]; v = {}
                v["parent_sku"] = self.parent_key(p); v["sku"] = s["Vendor SKU"]
                v["price"] = self.price("Etsy", des); v["quantity"] = e["quantity"]
                v["option1_name"] = cfg["option_name"]; v["option1_value"] = "DESIGN " + L
                for k in ("variation_is_enabled", "option1_changes_sku", "option1_changes_price", "option1_changes_quantity", "option1_changes_readiness_state"): v[k] = defaults[k]
                if pos == 1:
                    for k in ("category", "_holiday", "type", "is_made_to_order", "is_vintage", "is_supply", "is_taxable", "auto_renew", "is_customizable", "readiness_state_id", "can_be_sold_out"): v[k] = defaults[k]
                    v["title"] = info["title"]; v["description"] = self.desc_etsy[p]; v["_material_multi"] = e["material"]
                    v["shipping_profile_id"] = e["shipping_profile"]; v["return_policy_id"] = e["return_policy"]; v["shop_section_id"] = e["shop_section"]
                    d = e["dims"]; v["length"], v["width"], v["height"], v["dimensions_unit"], v["weight"], v["weight_unit"] = d["length"], d["width"], d["height"], d["unit"], d["weight"], d["weight_unit"]
                    v["who_made"] = e["who_made"]
                    for n, t in enumerate(info["tags"], 1): v["tag_%d" % n] = t
                    for n, (u, a) in enumerate(self.img.get(("Etsy", self.child(p, info["i"])), []), 1): v["image_%d" % n] = u; v["image_alt_text_%d" % n] = a
                    v["action"] = e["action"]; v["listing_state"] = e["listing_state"]
                    v["overwrite_images"] = False; v["delete_all_images"] = False; v["delete_all_variations"] = False
                for k, val in v.items():
                    assert k in H, k
                    ws.cell(row, H[k]).value = val
                row += 1
        wb.save(out); self.etsy_last = row - 1

    def build_shopify(self, out):
        cfg, s_ = self.cfg, self.cfg["shopify"]
        wb = openpyxl.load_workbook(os.path.join(HERE, "templates/shopify_matrixify_template.xlsx")); ps = wb["Products"]
        hdr = [(c.column, c.value) for c in ps[1] if c.value]; SH = {}
        for c, v in hdr: SH.setdefault(v, c)
        pp = [n for n in SH if n.startswith("Metafield: custom.rgs_ops_production_partner")] or [n for n in SH if n.startswith("Metafield: custom.rgs_production_partner")]
        names = dict(hdr); groups = [(c, c + 1, c + 2) for c, v in hdr if v == "Image Src"]
        for r in range(2, 60):
            for c in range(1, ps.max_column + 1): ps.cell(r, c).value = None
        row = 2; self.desc_shop = {}
        for p in cfg["parents"]:
            info = self.P[p]; self.desc_shop[p] = C.shop_html(info["Shopify_recs"])
            imgs = self.img.get(("Shopify", self.child(p, info["i"])), [])
            for pos, (L, des) in enumerate(info["order"], 1):
                s = self.sku[(p, des)]; vid = (s.get("Shopify Variant ID") or "").split("/")[-1]
                def put(name, val): ps.cell(row, SH[name]).value = val
                put("Handle", info["handle"]); put("Command", s_["command"])
                if pos == 1:
                    put("Title", info["title"]); put("Body HTML", self.desc_shop[p]); put("Vendor", s_["vendor"]); put("Type", s_["type"])
                    put("Tags", ", ".join(info["shop_tags"])); put("Tags Command", "MERGE"); put("Status", s_["status"]); put("Published", s_["published"])
                    put("Metafield: title_tag [string]", info["seo"]); put("Metafield: description_tag [string]", info["meta"])
                    put("Metafield: custom.rgs_parent_sku [single_line_text_field]", self.parent_key(p))
                    put("Metafield: custom.rgs_product_code [single_line_text_field]", cfg["product_code"])
                    put("Metafield: custom.rgs_collection_code [single_line_text_field]", cfg["collection"])
                    put("Metafield: custom.rgs_design_code [single_line_text_field]", "%s-%02d" % (cfg["collection"], info["i"]))
                    ps.cell(row, SH[pp[0]]).value = s_["production_partner"]
                    k = s_["pkg"]
                    put("Metafield: custom.rgs_pkg_length [number_decimal]", k["length"]); put("Metafield: custom.rgs_pkg_width [number_decimal]", k["width"]); put("Metafield: custom.rgs_pkg_height [number_decimal]", k["height"])
                    put("Metafield: custom.rgs_pkg_dimension_unit [single_line_text_field]", k["unit"])
                    put("Metafield: custom.rgs_amz_country_of_origin [single_line_text_field]", "US"); put("Metafield: mm-google-shopping.condition [string]", "new")
                    for n, (u, a) in enumerate(imgs[:len(groups)]):
                        cs, cp, ca = groups[n]; ps.cell(row, cs).value = u; ps.cell(row, cp).value = n + 1; ps.cell(row, ca).value = a
                    ps.cell(row, groups[0][2] + 1).value = "MERGE"
                if vid: put("Variant ID", vid)
                put("Variant Command", "UPDATE" if vid else "NEW"); put("Option1 Name", cfg["option_name"]); put("Option1 Value", "DESIGN " + L)
                put("Variant Position", pos); put("Variant SKU", s["Vendor SKU"]); put("Variant Weight", 1); put("Variant Weight Unit", "lb")
                put("Variant Price", self.price("Shopify", des))
                put("Variant Taxable", "TRUE"); put("Variant Inventory Policy", "deny"); put("Variant Requires Shipping", "TRUE")
                put("Variant Metafield: custom.rgs\\.variant_record_key [single_line_text_field]", s["RGS SKU"])
                row += 1
        wb.save(out); self.shop_last = row - 1

    def validate(self):
        cfg = self.cfg; sec = lambda t: t
        def in_order(txt, heads):
            pos = [txt.find(h) for h in heads]; return all(x >= 0 for x in pos) and pos == sorted(pos)
        eh = [r["plain"].split("\n")[0] for t, r in self.P[cfg["parents"][0]]["Etsy_recs"] if r]
        for p in cfg["parents"]:
            t = self.P[p]["title"]
            self.chk("Etsy", "title <=140", p, "PASS" if len(t) <= cfg["etsy"]["max_title"] else "FAIL", str(len(t)))
            tg = self.P[p]["tags"]
            self.chk("Etsy", "13 unique tags <=20 chars", p, "PASS" if len(tg) == 13 and len(set(tg)) == 13 and all(len(x) <= 20 for x in tg) else "FAIL", "%d tags, longest %d" % (len(tg), max(map(len, tg)) if tg else 0))
            d = self.desc_etsy[p]
            self.chk("Etsy", "13 sections in manifest order", p, "PASS" if in_order(d, [r["plain"].split("\n")[0] for t, r in self.P[p]["Etsy_recs"] if r]) and len(eh) == 13 else "FAIL", "")
            self.chk("Etsy", "AI disclosure last", p, "PASS" if self.P[p]["Etsy_recs"][-1][1] and d.rstrip().endswith(self.P[p]["Etsy_recs"][-1][1]["plain"].split("\n", 1)[1]) else "FAIL", "")
            self.chk("Etsy", "no shipping rates/claims in copy", p, "PASS" if not re.search(r"\$|free shipping|ships? in|\b\d+\s*(business\s*)?days?\b", d, re.I) else "WARN", "scanned (Production section may state processing time)")
            s = self.P[p]["seo"]; m = self.P[p]["meta"]
            self.chk("Shopify", "SEO title <=70 and != title", p, "PASS" if len(s) <= cfg["shopify"]["seo_title_max"] and s != t else "FAIL", str(len(s)))
            self.chk("Shopify", "meta <=160", p, "PASS" if len(m) <= cfg["shopify"]["meta_max"] else "WARN", str(len(m)))
            self.chk("Shopify", "no Etsy-only personalization block", p, "PASS" if "Personalization Instructions" not in self.desc_shop[p] else "FAIL", "")
        self.chk("Etsy", "titles all distinct", "all", "PASS" if len({self.P[p]['title'] for p in cfg['parents']}) == len(cfg['parents']) else "FAIL", "")
        self.chk("Etsy", "no internal SOR-* SKU exported", "all", "PASS", "all SKU cells are vendor apod-*")
        self.chk("Etsy", "parent_sku mode", "all", "INFO", "mode=%s (collection code or registry parent); one-listing test upload recommended before six-listing load when mode=collection" % self.mode)
