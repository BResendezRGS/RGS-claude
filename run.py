#!/usr/bin/env python3
"""python run.py --collection SOR-26-050 [--snapshot snapshots/SOR-26-050.json] [--out out] [--parent-sku-mode collection|registry]"""
import argparse, json, os, sys, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rgs.build import Run
import openpyxl
from openpyxl.styles import Font, PatternFill

ap = argparse.ArgumentParser()
ap.add_argument("--collection", required=True); ap.add_argument("--snapshot"); ap.add_argument("--out", default="out")
ap.add_argument("--skip-amazon", action="store_true"); ap.add_argument("--parent-sku-mode", choices=["collection", "registry"]); ap.add_argument("--allow-blockers", action="store_true")
a = ap.parse_args()
here = os.path.dirname(os.path.abspath(__file__))
cfg = json.load(open(os.path.join(here, "collections/%s.json" % a.collection)))
snap = json.load(open(a.snapshot or os.path.join(here, "snapshots/%s.json" % a.collection)))
mode = {"registry": "registry"}.get(a.parent_sku_mode, a.parent_sku_mode)
r = Run(cfg, snap, mode); os.makedirs(a.out, exist_ok=True)
ready = r.readiness()
seen = set()
for sev, item, scope, msg in r.log:
    key = (sev, msg if sev == "REVIEW" else (item, scope, msg))
    if key in seen: continue
    seen.add(key); print("%-8s %-22s %-16s %s" % (sev, item[:22], scope, msg))
if not ready and not a.allow_blockers:
    print("READINESS GATE FAILED - no files built (use --allow-blockers to build a dry run anyway)"); sys.exit(2)
e = os.path.join(a.out, "%s_Etsy_Bulk_Load.xlsx" % a.collection); s = os.path.join(a.out, "%s_Shopify_Matrixify.xlsx" % a.collection)
r.build_etsy(e); r.build_shopify(s); r.validate()
amz_out = None
if "amazon" in cfg and not a.skip_amazon:
    from rgs import amazon
    amz_out = os.path.join(a.out, "%s_Amazon_Template_Fill.xlsm" % a.collection); r.force = a.allow_blockers
    if not amazon.build(r, amz_out):
        print("AMAZON BLOCKERS - Amazon file not built"); amz_out = None
    for sev, item, scope, msg in r.log:
        if item.startswith(("Amazon", "Production", "Content Library")) and (sev, msg) not in seen: seen.add((sev, msg)); print("%-8s %-22s %-16s %s" % (sev, item[:22], scope, msg))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Run_Log"
ws.append(["Severity", "Item", "Scope", "Message"])
for l in r.log: ws.append(list(l))
w2 = wb.create_sheet("Checks"); w2.append(["Marketplace", "Check", "Scope", "Result", "Evidence"])
for m, rule, c, res, ev in r.checks: w2.append([m, rule, c, res, ev])
w3 = wb.create_sheet("Sources"); w3.append(["Section", "Parent", "Marketplace", "Library record", "Status", "Version"])
for p in cfg["parents"]:
    for mk in ("Etsy", "Shopify"):
        for t, rec in r.P[p][mk + "_recs"]: w3.append([t, p, mk, rec["id"] if rec else "MISSING", rec["status"] if rec else "", rec.get("ver") if rec else ""])
for w in (ws, w2, w3):
    for c in w[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
wb.save(os.path.join(a.out, "%s_Run_Report.xlsx" % a.collection))
cnt = collections.Counter(c[3] for c in r.checks)
print("built", e, s, amz_out, dict(cnt))
