"""Six-rotating-parent default test (rgs/rules.py, operator rule 2026-10-09).
Strips the structure keys from the SOR-26-004 config so rules.py must supply them, then builds Etsy and Shopify and checks:
six parents P01-P06 (no P07), fixed letters DESIGN A-G = designs 01-07 on every parent, identical apod SKU per design
across all six, DESIGN G last at its own price, one Shopify handle per parent, Amazon 1 parent + 7 children.
Also checks that a wrong parents list and SOR-26-050's explicit rotated letters behave as documented.
Run: python3 tests/test_rotating_default.py"""
import json, os, sys, tempfile
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); sys.path.insert(0, HERE)
import openpyxl
from rgs.build import Run
from rgs import amazon, rules

fail = []
def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg); cond or fail.append(msg)

COLL = "SOR-26-004"
snap = json.load(open(os.path.join(HERE, "snapshots/%s.json" % COLL)))
cfg = json.load(open(os.path.join(HERE, "collections/%s.json" % COLL)))
for k in ("listing_structure", "listing_structure_note", "letter_mapping", "parents"):
    cfg.pop(k, None)
r = Run(cfg, snap); r.readiness(); r.force = True   # images are pending for SOR-26-004, same as --allow-blockers
P = ["%s-P%02d" % (COLL, n) for n in range(1, 7)]
check(cfg["listing_structure"] == "rotating_parents" and cfg["letter_mapping"] == "fixed", "defaults: rotating_parents with fixed letters")
check(cfg["parents"] == P, "defaults: parents P01-P06, no parent for ready-made design")
check(not any(m[0] == "BLOCKER" and m[1] == "Rules" for m in r.log), "no Rules blocker on defaults")
check(all(r.P[p]["order"] == list(zip("ABCDEFG", range(1, 8))) for p in P), "every parent: DESIGN A-G = designs 01-07")
check([r.P[p]["i"] for p in P] == [1, 2, 3, 4, 5, 6], "parent N leads with design N")

d = tempfile.mkdtemp(); e, s, a = (os.path.join(d, n) for n in ("e.xlsx", "s.xlsx", "a.xlsm"))
r.build_etsy(e); r.build_shopify(s); r.validate(); amz = amazon.build(r, a)
ws = openpyxl.load_workbook(e)["Template"]; H = {c.value: c.column for c in ws[1] if c.value}
rows = [[ws.cell(i, H[h]).value for h in ("sku", "option1_value", "price", "title")] for i in range(2, ws.max_row + 1) if ws.cell(i, H["sku"]).value]
check(len(rows) == 42, "Etsy: 6 listings x 7 rows (%d)" % len(rows))
by_letter = {}
for sku, opt, price, _ in rows: by_letter.setdefault(opt, set()).add(sku)
check(all(len(v) == 1 for v in by_letter.values()) and len({next(iter(v)) for v in by_letter.values()}) == 7, "Etsy: one apod SKU per design letter, shared by all six")
check(all(str(x[0]).startswith("apod-") for x in rows), "Etsy: only apod-* SKUs exported")
check(all(rows[k * 7 + 6][1] == "DESIGN G" and rows[k * 7 + 6][2] == 9.99 and rows[k * 7][2] == 14.98 for k in range(6)), "Etsy: DESIGN G last at 9.99 on each listing, A-F 14.98")
check(len({x[3] for x in rows if x[3]}) == 6, "Etsy: six distinct titles")
ps = openpyxl.load_workbook(s)["Products"]
hc = next(c.column for c in ps[1] if c.value == "Handle")
check(len({ps.cell(i, hc).value for i in range(2, ps.max_row + 1) if ps.cell(i, hc).value}) == 6, "Shopify: six product handles")
check(amz and len(r.amz_rows) == 8, "Amazon: 1 parent + 7 children")

bad = json.load(open(os.path.join(HERE, "collections/%s.json" % COLL))); bad["parents"] = P + ["%s-P07" % COLL]
log = []; rules.apply(bad, log)
check(any(m[0] == "BLOCKER" and "rotating_parents requires" in m[3] for m in log), "a -P07 parent (ready-made lead) is a BLOCKER")

leg = json.load(open(os.path.join(HERE, "collections/SOR-26-050.json"))); rules.apply(leg, [])
check(leg["letter_mapping"] == "rotated" and rules.variant_order(leg, 2)[0] == ("A", 2) and rules.variant_order(leg, 2)[-1] == ("G", 7), "SOR-26-050 keeps rotated letters, G still last")
print("\n%d failure(s)" % len(fail)); sys.exit(1 if fail else 0)
