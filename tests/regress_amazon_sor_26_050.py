"""Regression: rebuilt Amazon fill vs the submitted test file (v2). Expected difference: Design 07 description opening only."""
import openpyxl, warnings; warnings.filterwarnings("ignore")
from openpyxl.utils import get_column_letter as L
OLD = "/tmp/claude-0/-home-claude/433e5365-a398-59b7-87cb-cf8dd3054a5e/scratchpad/SOR-26-050_Amazon_Template_Fill_v2.xlsm"
NEW = "out/SOR-26-050_Amazon_Template_Fill.xlsm"
def rows(p):
    ws = openpyxl.load_workbook(p, keep_vba=True)["Template"]
    return {(r, L(c)): str(ws.cell(r, c).value) for r in range(8, 16) for c in range(1, 485) if ws.cell(r, c).value not in (None, "")}
a, b = rows(OLD), rows(NEW)
d = [(k, a.get(k), b.get(k)) for k in sorted(set(a) | set(b)) if a.get(k) != b.get(k)]
print("cells compared", len(set(a) | set(b)), "differing", len(d))
for k, x, y in d: print(k, "\n  old:", (x or "")[:110], "\n  new:", (y or "")[:110])
