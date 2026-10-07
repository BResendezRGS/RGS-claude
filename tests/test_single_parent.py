"""Single-parent rule test (rgs/rules.py, 2026-10-07).
Builds a one-parent fixture from the SOR-26-050 snapshot, shaped like the ORN-26-051 registry
(parent row RGS SKU == collection code, children <COLL>-01..07), then runs Etsy, Shopify and Amazon builders.
Run: python3 tests/test_single_parent.py"""
import copy, json, os, sys, tempfile
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); sys.path.insert(0, HERE)
import openpyxl
from rgs.build import Run
from rgs import amazon

COLL = "SOR-26-050"
snap = json.load(open(os.path.join(HERE, "snapshots/%s.json" % COLL)))
cfg = json.load(open(os.path.join(HERE, "collections/%s.json" % COLL)))
for k in ("listing_structure", "listing_structure_note", "parents", "production_partner", "ready_made_design"):
    cfg.pop(k, None)                                   # let rules.py supply the defaults
cfg["shopify"].pop("production_partner", None)

T = copy.deepcopy(snap["tables"]); src = "-P01"
def child(s): return s.replace(src, "") if isinstance(s, str) else s
sku = [dict(r, **{"RGS SKU": child(r["RGS SKU"]), "RGS Parent SKU": COLL}) for r in T["SKU Registries"] if r["RGS Parent SKU"] == COLL + src]
sku.append({"RGS SKU": COLL, "RGS Parent SKU": COLL, "Vendor SKU": "PENDING_VENDOR_SKU"})   # registry parent row
st = [dict(r, **{"RGS Child SKU": child(r["RGS Child SKU"]), "RGS Parent SKU": COLL}) for r in T["Listing Build Staging"] if r.get("RGS Parent SKU") == COLL + src]
img = [dict(r, **{"Product SKU": child(r["Product SKU"])}) for r in T["Image URL Console"] if str(r.get("Product SKU", "")).endswith(src)]
T.update({"SKU Registries": sku, "Listing Build Staging": st, "Image URL Console": img})
fx = {"collection": COLL, "tables": T}

r = Run(cfg, fx); ok = r.readiness()
fail = []
def check(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg); cond or fail.append(msg)

check(cfg["listing_structure"] == "single_parent" and cfg["parents"] == [COLL], "defaults: single parent keyed to collection code")
check(cfg["production_partner"] == "ICP" and cfg["shopify"]["production_partner"] == "InnerCircle Prints", "defaults: partner ICP -> InnerCircle Prints")
check(cfg["ready_made_design"] == 7, "defaults: ready-made design 07")
check(ok, "readiness gate passes: " + "; ".join(m for s, _, _, m in r.log if s == "BLOCKER"))
check((COLL, 0) not in r.sku and len(r.sku) == 7, "registry parent row skipped, 7 design rows keyed")
check([(L, d) for L, d in r.P[COLL]["order"]] == list(zip("ABCDEFG", range(1, 8))), "variant order DESIGN A..G = designs 01..07")

d = tempfile.mkdtemp(); e, s, a = (os.path.join(d, n) for n in ("e.xlsx", "s.xlsx", "a.xlsm"))
r.build_etsy(e); r.build_shopify(s); r.validate(); amz = amazon.build(r, a)
ws = openpyxl.load_workbook(e)["Template"]; H = {c.value: c.column for c in ws[1] if c.value}
rows = [[ws.cell(i, H[h]).value for h in ("parent_sku", "sku", "option1_value", "price", "title")] for i in range(2, ws.max_row + 1) if ws.cell(i, H["sku"]).value]
check(len(rows) == 7, "Etsy: one listing, 7 variation rows (%d)" % len(rows))
check(all(x[0] == COLL for x in rows), "Etsy: parent_sku = collection code on every row")
check(all(str(x[1]).startswith("apod-") for x in rows), "Etsy: only vendor apod-* SKUs exported")
check(rows[-1][2] == "DESIGN G" and rows[-1][3] == 9.99 and rows[0][3] == 14.98, "Etsy: DESIGN G last at 9.99, A-F 14.98")
check(sum(1 for x in rows if x[4]) == 1, "Etsy: listing-level fields on first row only")
ps = openpyxl.load_workbook(s)["Products"]
hc = next(c.column for c in ps[1] if c.value == "Handle")
hs = {ps.cell(i, hc).value for i in range(2, ps.max_row + 1) if ps.cell(i, hc).value}
check(len(hs) == 1, "Shopify: one product handle (%s)" % hs)
check(amz and len(r.amz_rows) == 8, "Amazon: 1 parent + 7 children")
check(not any(c[3] == "FAIL" for c in r.checks), "no FAIL checks: " + ", ".join("%s %s" % (c[1], c[2]) for c in r.checks if c[3] == "FAIL"))
print("\n%d failure(s)" % len(fail)); sys.exit(1 if fail else 0)
